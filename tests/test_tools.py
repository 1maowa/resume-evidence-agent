from __future__ import annotations

import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from docx import Document

from career_match_agent.models import (
    JobPostingRequest,
    LLMAnalysisRequest,
    ParseResumeRequest,
    RequirementAssessment,
    SaveReportRequest,
)
from career_match_agent.runtime import ToolRuntime
from career_match_agent.tools import (
    ToolError,
    fetch_job_posting,
    parse_resume,
    score_llm_analysis,
)


ROOT = Path(__file__).resolve().parents[1]
RESUME = ROOT / "samples" / "resume_demo.md"
JOB = ROOT / "samples" / "job_description_demo.md"


class HtmlHandler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:  # noqa: N802
        if self.path == "/challenge":
            body = (
                "<html><head><title>请稍候 - BOSS直聘</title></head>"
                "<body><p>正在进行安全验证，请启用 JavaScript 和 Cookie。</p></body></html>"
            ).encode("utf-8")
        else:
            body = (
                "<html><head><title>Python 开发实习生</title></head>"
                "<body><h1>岗位职责</h1><p>使用 Python、SQL 和 Docker。</p>"
                "<script>ignore me</script></body></html>"
            ).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


class ToolTests(unittest.TestCase):
    def test_parse_resume_redacts_pii(self) -> None:
        result = parse_resume(ParseResumeRequest(path=str(RESUME), redact_pii=True))
        self.assertIn("[EMAIL_REDACTED]", result.text)
        self.assertIn("[PHONE_REDACTED]", result.text)
        self.assertNotIn("linzhou@example.com", result.text)
        self.assertNotIn("138-0000-1234", result.text)
        self.assertTrue(result.pii_redacted)

    def test_parse_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "resume.docx"
            document = Document()
            document.add_paragraph("Python 与 SQL 项目经历")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "Pandas 数据清洗"
            document.save(path)

            result = parse_resume(ParseResumeRequest(path=str(path), redact_pii=False))
            self.assertIn("Python 与 SQL", result.text)
            self.assertIn("Pandas 数据清洗", result.text)
            self.assertEqual(result.file_type, "docx")

    def test_rejects_scanned_empty_pdf_like_text(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "empty.txt"
            path.write_text("short", encoding="utf-8")
            with self.assertRaises(ToolError):
                parse_resume(ParseResumeRequest(path=str(path)))

    def test_llm_analysis_scores_weighted_requirements(self) -> None:
        request = LLMAnalysisRequest(
            job_title="AI Infra 实习生",
            resume_summary=(
                "候选人的主要方向是大模型推理与训练系统，具有 CUDA 算子、"
                "vLLM 推理优化和分布式训练经验，但缺少大规模集群运维证据。"
            ),
            strengths=["vLLM 推理优化", "CUDA 与 Triton 算子", "分布式训练"],
            requirements=[
                _requirement("CUDA 编程", "required", True, "实现 GEMM 与 FlashAttention"),
                _requirement("vLLM 推理优化", "required", True, "基于 vLLM 部署 Llama-3-8B"),
                _requirement("Kubernetes", "required", True, "基础设施：Kubernetes"),
                _requirement("RDMA", "preferred", False),
                _requirement("千卡集群经验", "preferred", False),
            ],
            recommendations=[
                "补充 RDMA 或高性能通信库实践。",
                "明确说明集群规模、吞吐和延迟指标。",
            ],
        )
        analysis = score_llm_analysis(request)
        self.assertEqual(analysis.score, 82)
        self.assertEqual(analysis.matched_count, 3)
        self.assertIn("RDMA", analysis.missing_requirements)

    def test_matched_requirement_requires_resume_evidence(self) -> None:
        with self.assertRaises(ValueError):
            RequirementAssessment(
                requirement="CUDA 编程",
                category="technical",
                priority="required",
                matched=True,
                evidence="",
            )

    def test_detects_job_challenge_page(self) -> None:
        server = ThreadingHTTPServer(("127.0.0.1", 0), HtmlHandler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            url = f"http://127.0.0.1:{server.server_port}/challenge"
            with self.assertRaises(ToolError) as context:
                fetch_job_posting(JobPostingRequest(source=url))
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=3)
        self.assertIn("粘贴正文", str(context.exception))

    def test_reads_static_job_file(self) -> None:
        job = fetch_job_posting(JobPostingRequest(source=str(JOB)))
        self.assertEqual(job.title, "数据分析实习生")
        self.assertIn("Pandas", job.text)

    def test_fetches_remote_html(self) -> None:
        server = ThreadingHTTPServer(("127.0.0.1", 0), HtmlHandler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            url = f"http://127.0.0.1:{server.server_port}/job"
            result = fetch_job_posting(JobPostingRequest(source=url))
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=3)
        self.assertEqual(result.title, "Python 开发实习生")
        self.assertIn("使用 Python、SQL 和 Docker", result.text)
        self.assertNotIn("ignore me", result.text)

    def test_save_report_is_confined_to_project(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            runtime = ToolRuntime(Path(temp_dir))
            saved = runtime.save_report(
                SaveReportRequest(
                    markdown="# 测试报告\n\n这是用于验证报告保存流程的正文内容。",
                    output_path="reports/test.md",
                )
            )
            self.assertTrue(Path(saved["output_path"]).is_file())
            with self.assertRaises(ValueError):
                runtime.save_report(
                    SaveReportRequest(
                        markdown="# 越界报告\n\n这段内容不应被写到项目目录之外。",
                        output_path=str(ROOT / "escaped.md"),
                    )
                )


def _requirement(
    requirement: str,
    priority: str,
    matched: bool,
    evidence: str | None = None,
) -> RequirementAssessment:
    return RequirementAssessment(
        requirement=requirement,
        category="technical",
        priority=priority,
        matched=matched,
        evidence=evidence,
    )

if __name__ == "__main__":
    unittest.main()

