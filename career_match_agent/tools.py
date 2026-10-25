from __future__ import annotations

import re
import unicodedata
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from .models import (
    EvidenceItem,
    JobPosting,
    JobPostingRequest,
    LLMAnalysisRequest,
    MatchAnalysis,
    ParseResumeRequest,
    ResumeDocument,
)


MAX_DOCUMENT_CHARS = 60_000
MAX_DOWNLOAD_BYTES = 2_000_000
SUPPORTED_RESUME_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}
SUPPORTED_JOB_SUFFIXES = {".txt", ".md", ".markdown", ".html", ".htm"}
JOB_CHALLENGE_MARKERS = (
    "请稍候",
    "安全验证",
    "访问验证",
    "滑动验证",
    "验证码",
    "访问过于频繁",
    "just a moment",
    "attention required",
    "enable javascript and cookies",
    "access denied",
)


class ToolError(ValueError):
    """An expected, user-correctable tool error."""


class _VisibleHTMLParser(HTMLParser):
    BLOCK_TAGS = {
        "article",
        "br",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "li",
        "p",
        "section",
        "title",
        "tr",
    }
    IGNORED_TAGS = {"script", "style", "noscript", "svg"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._ignored_depth = 0
        self._parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self.IGNORED_TAGS:
            self._ignored_depth += 1
        elif tag in self.BLOCK_TAGS:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.IGNORED_TAGS and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag in self.BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self._parts.append(data)

    def text(self) -> str:
        return "".join(self._parts)


def clean_text(value: str, *, preserve_lines: bool = True) -> str:
    value = unicodedata.normalize("NFKC", value).replace("\x00", " ")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    if preserve_lines:
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
        return "\n".join(line for line in lines if line)
    return re.sub(r"\s+", " ", value).strip()


def truncate_document(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_DOCUMENT_CHARS:
        return text, False
    return text[:MAX_DOCUMENT_CHARS], True


def redact_pii(text: str) -> str:
    text = re.sub(
        r"(?<![\w.+-])[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?!\w)",
        "[EMAIL_REDACTED]",
        text,
    )
    text = re.sub(
        r"(?<!\d)(?:\+?86[-\s]?)?1[3-9]\d(?:[-\s]?\d{4}){2}(?!\d)",
        "[PHONE_REDACTED]",
        text,
    )
    text = re.sub(
        r"(?<!\d)\+?[1-9]\d{0,2}(?:[-.\s]\d{2,4}){2,3}(?!\d)",
        "[PHONE_REDACTED]",
        text,
    )
    return text


def parse_resume(request: ParseResumeRequest) -> ResumeDocument:
    path = Path(request.path).expanduser().resolve()
    if not path.is_file():
        raise ToolError(f"简历文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_RESUME_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_RESUME_SUFFIXES))
        raise ToolError(f"不支持的简历格式 {suffix}，支持：{supported}")

    if suffix == ".pdf":
        text = _read_pdf(path)
    elif suffix == ".docx":
        text = _read_docx(path)
    else:
        text = path.read_text(encoding="utf-8-sig")

    text = clean_text(text)
    if len(text) < 20:
        raise ToolError("简历中没有提取到足够的可读文本。若是扫描 PDF，请先进行 OCR。")
    if request.redact_pii:
        text = redact_pii(text)
    text, truncated = truncate_document(text)
    return ResumeDocument(
        source=str(path),
        file_type=suffix.lstrip("."),
        text=text,
        character_count=len(text),
        truncated=truncated,
        pii_redacted=request.redact_pii,
    )


def fetch_job_posting(request: JobPostingRequest) -> JobPosting:
    source = request.source.strip()
    if not source:
        raise ToolError("岗位来源不能为空。")
    is_windows_path = bool(re.match(r"^[A-Za-z]:[\\/]", source))
    parsed = urlparse(source)
    if parsed.scheme.lower() in {"http", "https"}:
        raw_text = _download_url(source, request.timeout_seconds)
        source_type = "url"
        title = _extract_html_title(raw_text) or _guess_job_title(raw_text)
        text = _html_to_text(raw_text)
        if _looks_like_challenge_page(raw_text, title, text):
            raise ToolError(
                "招聘网站返回了验证或反爬页面，未取得真实岗位正文。"
                "请复制岗位职责和任职要求，选择“粘贴正文”后重新分析。"
            )
    else:
        if parsed.scheme and parsed.scheme.lower() != "file":
            if not is_windows_path:
                raise ToolError("岗位来源只支持 HTTP/HTTPS 地址或本地文件。")
        if parsed.scheme.lower() == "file":
            local_source = source.removeprefix("file://")
        else:
            local_source = source
        path = Path(local_source).expanduser().resolve()
        if not path.is_file():
            raise ToolError(f"岗位文件不存在：{path}")
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_JOB_SUFFIXES:
            supported = ", ".join(sorted(SUPPORTED_JOB_SUFFIXES))
            raise ToolError(f"不支持的岗位文件格式 {suffix}，支持：{supported}")
        raw_text = path.read_text(encoding="utf-8-sig", errors="replace")
        source_type = "file"
        title = _extract_html_title(raw_text) or _guess_job_title(raw_text)
        text = _html_to_text(raw_text) if suffix in {".html", ".htm"} else raw_text

    text = clean_text(text)
    if len(text) < 20:
        raise ToolError("岗位描述中没有提取到足够的可读文本。")
    text, truncated = truncate_document(text)
    return JobPosting(
        source=source,
        source_type=source_type,
        title=title,
        text=text,
        character_count=len(text),
        truncated=truncated,
    )


def score_llm_analysis(request: LLMAnalysisRequest) -> MatchAnalysis:
    priority_weights = {"required": 3, "preferred": 1}
    total_weight = sum(priority_weights[item.priority] for item in request.requirements)
    matched_weight = sum(
        priority_weights[item.priority]
        for item in request.requirements
        if item.matched
    )
    score = round(100 * matched_weight / total_weight)
    evidence = [
        EvidenceItem(
            requirement=item.requirement,
            category=item.category,
            priority=item.priority,
            matched=item.matched,
            evidence=clean_text(item.evidence, preserve_lines=False) if item.evidence else None,
        )
        for item in request.requirements
    ]
    missing = [item.requirement for item in request.requirements if not item.matched]
    required_count = sum(
        1 for item in request.requirements if item.priority == "required"
    )
    preferred_count = len(request.requirements) - required_count
    return MatchAnalysis(
        job_title=request.job_title.strip(),
        score=score,
        level=_score_level(score),
        total_requirements=len(request.requirements),
        matched_count=sum(1 for item in request.requirements if item.matched),
        missing_requirements=missing,
        evidence=evidence,
        resume_summary=request.resume_summary.strip(),
        strengths=request.strengths,
        recommendations=request.recommendations,
        notes=[
            "岗位要求、简历证据和匹配判断由 LLM 生成，工具负责结构校验与加权计分。",
            f"共拆解 {len(request.requirements)} 项要求，其中核心要求 {required_count} 项、"
            f"加分项 {preferred_count} 项。",
            "核心要求权重为 3，加分项权重为 1；得分等于已匹配权重除以总权重。",
            "文本证据不等于能力认证，面试前仍需核对事实。",
        ],
    )


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ToolError("解析 PDF 需要安装 pypdf。") from exc
    try:
        reader = PdfReader(str(path))
        pages: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                pages.append(f"[第{index}页]\n{page_text}")
        return "\n\n".join(pages)
    except Exception as exc:
        raise ToolError(f"PDF 解析失败：{exc}") from exc


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("解析 DOCX 需要安装 python-docx。") from exc
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ToolError(f"DOCX 解析失败：{exc}") from exc
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _download_url(url: str, timeout_seconds: int) -> str:
    request = Request(
        url,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 Chrome/124 Safari/537.36"
            )
        },
    )
    try:
        with urlopen(request, timeout=timeout_seconds) as response:
            payload = response.read(MAX_DOWNLOAD_BYTES + 1)
            charset = response.headers.get_content_charset() or "utf-8"
    except Exception as exc:
        raise ToolError(f"岗位页面抓取失败：{exc}") from exc
    if len(payload) > MAX_DOWNLOAD_BYTES:
        raise ToolError("岗位页面超过 2 MB，已停止下载。请改用本地岗位文本。")
    return payload.decode(charset, errors="replace")


def _html_to_text(value: str) -> str:
    parser = _VisibleHTMLParser()
    parser.feed(value)
    parser.close()
    return parser.text()


def _extract_html_title(value: str) -> str:
    match = re.search(r"<title[^>]*>(.*?)</title>", value, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return ""
    title = re.sub(r"<[^>]+>", " ", match.group(1))
    return clean_text(title, preserve_lines=False)[:120]


def _guess_job_title(value: str) -> str:
    for line in value.splitlines():
        cleaned = re.sub(r"^\s*[#*\-\d.)、]+\s*", "", line).strip()
        if 2 <= len(cleaned) <= 80 and not cleaned.lower().startswith(("http", "岗位职责")):
            return cleaned
    return "岗位描述"


def _score_level(score: int) -> str:
    if score >= 80:
        return "高匹配"
    if score >= 60:
        return "中等匹配"
    return "低匹配"


def _looks_like_challenge_page(raw_html: str, title: str, text: str) -> bool:
    lowered_title = title.casefold()
    lowered_text = text[:5000].casefold()
    lowered_html = raw_html[:20000].casefold()
    title_hit = any(marker.casefold() in lowered_title for marker in JOB_CHALLENGE_MARKERS)
    text_hit = any(marker.casefold() in lowered_text for marker in JOB_CHALLENGE_MARKERS)
    html_hit = any(marker.casefold() in lowered_html for marker in JOB_CHALLENGE_MARKERS)
    return title_hit or text_hit or (html_hit and len(text) < 1500)

