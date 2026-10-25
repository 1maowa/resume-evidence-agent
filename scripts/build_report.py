from __future__ import annotations

import argparse
import hashlib
import os
import shutil
from zipfile import ZIP_DEFLATED, ZipFile
from pathlib import Path

from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = PROJECT_ROOT / "实验1-实验报告模版.docx"
DEFAULT_OUTPUT = PROJECT_ROOT / "实验1-实验报告-完整版.docx"
ARCHITECTURE_IMAGE = PROJECT_ROOT / "assets" / "architecture.png"
TOOL_FLOW_IMAGE = PROJECT_ROOT / "assets" / "tool_flow.png"
WEB_SCREENSHOT = PROJECT_ROOT / "assets" / "web_workspace.png"
TRACE_IMAGE = PROJECT_ROOT / "assets" / "llm_tool_trace.png"
MCP_IMAGE = PROJECT_ROOT / "assets" / "mcp_tools.png"
TEMPLATE_SHA256 = "0a57ce9f6102fd58e2eed2b82560cae8331a2075e7642c687a5995558b50e924"
FONT_NAME = "微软雅黑"
INK = "000000"
MUTED = "5F6368"
BORDER = "D9D9D9"
PLACEHOLDER_FILL = "F4F6F8"


SECTION_4 = [
    (
        "该 Agent 面向 AI 岗位求职场景，将简历与目标岗位做语义化证据匹配。输入支持 "
        "PDF、DOCX、TXT/Markdown 简历及招聘 URL 或岗位文件。四项工具分别是："
        "parse_resume 解析文本并脱敏；fetch_job_posting 读取岗位并识别验证页；"
        "record_analysis 接收 LLM 拆出的技术、经验、教育和软性要求，校验已匹配项必须"
        "含简历原文，按核心要求权重 3、加分项权重 1 计分，并生成简历速览；"
        "save_report 保存 Markdown 报告。"
    ),
    (
        "FastAPI Web 工作台支持拖放、粘贴和文件上传，直接展示分数、缺口与逐项证据。"
        "上下文通过 OpenAI 兼容 function calling 和 FastMCP stdio MCP Server 接入，"
        "二者复用 ToolRuntime 与 Pydantic 模型。LLM 负责语义理解，工具负责校验、计分和"
        "落盘；外部文本按不可信数据处理。"
    ),
]

SECTION_5 = [
    (
        "系统只提供 LLM 模式，岗位要求拆解和简历证据判断均由模型完成。执行 python -m "
        "career_match_agent.web_app --port 8000 后，浏览器通过 FastAPI 上传材料；Agent "
        "依次调用 parse_resume、fetch_job_posting、record_analysis、save_report。模型"
        "拆出独立要求，逐条给出 matched、priority、category 和简历原文 evidence；若标记"
        "已匹配但缺少证据，Pydantic 会拒绝并触发重试。"
    ),
    (
        "真实 DeepSeek 示例中，LLM 从 AI Infra 岗位拆出 35 项要求，29 项找到原文证据，"
        "加权得分 85，并生成简历速览、核心优势和 5 条修改建议。17 项自动测试覆盖反爬"
        "识别、证据校验、加权计分、Web 接口、函数调用和 MCP；浏览器已在 1440×1000、"
        "390×844 两种尺寸检查。图 4 为实际 function calling 轨迹，图 5 为 MCP 返回。"
    ),
]

SECTION_6 = [
    (
        "本实验使用 Codex 进行 vibe coding。先由我确定目标、数据边界和提示词规则，"
        "AI 生成 Pydantic 模型、解析器、OpenAI 函数调用循环、FastMCP Server、"
        "FastAPI 接口以及 HTML/CSS/JavaScript 工作台和测试骨架。"
    ),
    (
        "初始版本用固定技能表，AI Infra 岗位只识别出 3 项并错误得到 100 分。后来改为"
        "由 LLM 拆解完整要求、逐项寻找简历原文，并调用 record_analysis 校验证据与加权"
        "计分；已匹配项缺少证据会被拒绝。另增加简历速览与建议，并识别 BOSS“请稍候”"
        "反爬页。Windows 路径误判、电话号码脱敏、MCP 结构化返回、移动端标题断行和"
        "结果定位问题均通过测试与真实浏览器检查修正。"
    ),
    (
        "第二个问题是 LLM 输出必须可校验。record_analysis 的参数包含多项岗位要求、"
        "优先级、类别和长段证据，实际运行时曾出现一次工具参数 JSON 被截断。程序没有"
        "接受不完整结果，而是把解析错误回传给模型重试；第二次调用成功后生成了结构化"
        "分析。由此确认了使用 Pydantic schema、限制最大步骤、统一错误回传和保留工具"
        "轨迹的必要性，也避免了模型用自由文本直接覆盖分数。"
    ),
    (
        "第三个问题是输入可信度与隐私。直接抓取 BOSS 页面时会得到“请稍候”验证页，"
        "系统现在识别该类页面并提示粘贴正文；上传简历默认替换邮箱和电话，API Key 只从"
        "本地 .env 读取，不进入浏览器或报告。前端通过 Playwright 在 1440×1000 和"
        "390×844 两种窗口检查，修复了移动端标题断行、favicon 404、结果自动定位和"
        "表单可访问性问题。"
    ),
    (
        "最终采用“生成、测试、修正”的闭环。每次结构变更后执行 17 项自动测试，并用"
        "真实 DeepSeek API 对 PDF 简历和 AI Infra 岗位做端到端验证，生成 35 项要求、"
        "29 项证据、85 分及完整 Markdown 报告。人工主要负责主题范围、数据边界、提示词"
        "约束和最终事实判断。"
    ),
]


def main() -> None:
    parser = argparse.ArgumentParser(description="从课程模板生成已填写的实验报告。")
    parser.add_argument(
        "--template",
        default=str(DEFAULT_TEMPLATE),
        help="保留不改的原始 DOCX 模板路径。",
    )
    parser.add_argument(
        "--out",
        default=str(DEFAULT_OUTPUT),
        help="最终实验报告 DOCX 路径，不能与模板相同。",
    )
    args = parser.parse_args()

    template = Path(args.template).resolve()
    output = Path(args.out).resolve()
    _verify_source(template, output)
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)

    document = Document(output)
    if len(document.paragraphs) != 27:
        raise RuntimeError(
            f"模板段落结构已变化，预期 27 段，实际 {len(document.paragraphs)} 段。"
        )
    paragraphs = document.paragraphs

    _replace_paragraph(
        paragraphs[2],
        "1. 姓名：【待填写：姓名】    2. 学号：【待填写：学号】",
    )
    _replace_paragraph(
        paragraphs[5],
        "【待填写：GitHub 仓库地址（包含 prompts/system_prompt.md）】",
    )

    _fill_paragraph(paragraphs[9], SECTION_4[0], space_after=4)
    _fill_paragraph(paragraphs[10], SECTION_4[1], space_after=6)
    _remove_paragraph(paragraphs[11])
    _remove_paragraph(paragraphs[12])

    page_break_paragraph = paragraphs[13]
    _insert_figure_before(page_break_paragraph)

    _fill_paragraph(paragraphs[15], SECTION_5[0], space_after=4)
    _fill_paragraph(paragraphs[16], SECTION_5[1], space_after=8)
    _remove_paragraph(paragraphs[17])

    anchor = paragraphs[18]
    _insert_picture_before(
        anchor,
        WEB_SCREENSHOT,
        caption="图 3  Web 工作台中的 LLM 分析与简历速览",
        width=Inches(6.25),
        alt_text="Web 工作台显示 AI Infra 岗位匹配分数、岗位要求和简历速览",
    )
    _insert_picture_before(
        anchor,
        TRACE_IMAGE,
        caption="图 4  LLM 函数调用与实际工具轨迹",
        width=Inches(6.25),
        alt_text="LLM 依次调用简历解析、岗位读取、结构化分析和报告保存工具",
    )
    _insert_picture_before(
        anchor,
        MCP_IMAGE,
        caption="图 5  MCP tools/list 返回的四个工具",
        width=Inches(6.25),
        alt_text="MCP 客户端通过 tools/list 发现四个工具及其参数",
    )

    for paragraph in paragraphs[18:26]:
        _remove_paragraph(paragraph)

    paragraphs[26].paragraph_format.keep_together = True
    paragraphs[26].paragraph_format.keep_with_next = True
    for text in SECTION_6:
        paragraph = document.add_paragraph()
        _fill_paragraph(paragraph, text, space_after=5)

    document.save(output)
    _restore_preserved_parts(template, output)

    _verify_source(template, output)
    print(output)


def _verify_source(template: Path, output: Path) -> None:
    if not template.is_file():
        raise FileNotFoundError(f"模板不存在：{template}")
    if template == output:
        raise ValueError("输出路径不能与保留模板相同。")
    digest = hashlib.sha256(template.read_bytes()).hexdigest()
    if digest != TEMPLATE_SHA256:
        raise RuntimeError("模板 SHA-256 与蒸馏记录不一致，请重新检查模板。")


def _replace_paragraph(paragraph, text: str) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    _add_run(paragraph, text, size=12)


def _fill_paragraph(paragraph, text: str, *, space_after: float = 4) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.12
    _add_run(paragraph, text, size=12)


def _add_run(
    paragraph,
    text: str,
    *,
    size: float,
    bold: bool = False,
    color: str = INK,
):
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _insert_figure_before(anchor) -> None:
    picture_paragraph = anchor.insert_paragraph_before()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(2)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.keep_with_next = True
    run = picture_paragraph.add_run()
    inline_shape = run.add_picture(str(ARCHITECTURE_IMAGE), width=Inches(6.25))
    inline_shape._inline.docPr.set("descr", "简历与岗位证据匹配 Agent 系统架构图")
    inline_shape._inline.docPr.set("title", "架构图")

    caption = anchor.insert_paragraph_before()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(3)
    _add_run(caption, "图 1  简历与岗位证据匹配 Agent 架构", size=10.5, color=MUTED)

    flow_paragraph = anchor.insert_paragraph_before()
    flow_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow_paragraph.paragraph_format.space_before = Pt(8)
    flow_paragraph.paragraph_format.space_after = Pt(3)
    flow_paragraph.paragraph_format.keep_with_next = True
    flow_run = flow_paragraph.add_run()
    flow_shape = flow_run.add_picture(str(TOOL_FLOW_IMAGE), width=Inches(6.25))
    flow_shape._inline.docPr.set("descr", "Agent 工具调用循环与示例运行结果图")
    flow_shape._inline.docPr.set("title", "工具调用循环")

    flow_caption = anchor.insert_paragraph_before()
    flow_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow_caption.paragraph_format.space_after = Pt(3)
    _add_run(flow_caption, "图 2  工具调用循环与示例运行结果", size=10.5, color=MUTED)


def _insert_picture_before(
    anchor,
    image_path: Path,
    *,
    caption: str,
    width,
    alt_text: str,
) -> None:
    picture_paragraph = anchor.insert_paragraph_before()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(8)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_run = picture_paragraph.add_run()
    picture_shape = picture_run.add_picture(str(image_path), width=width)
    picture_shape._inline.docPr.set("descr", alt_text)
    picture_shape._inline.docPr.set("title", caption)

    caption_paragraph = anchor.insert_paragraph_before()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(4)
    _add_run(caption_paragraph, caption, size=10.5, color=MUTED)


def _insert_screenshot_placeholder(
    document: Document,
    anchor,
    *,
    title: str,
    hint: str,
    caption: str,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(6.55)
    cell = table.cell(0, 0)
    cell.width = Inches(6.55)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell, top=100, start=120, bottom=100, end=120)
    _set_table_borders(table)
    _shade_cell(cell, PLACEHOLDER_FILL)
    row = table.rows[0]
    row.height = Inches(1.15)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    _mark_row_as_header(row)

    label = cell.paragraphs[0]
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(2)
    _add_run(label, title, size=11.5, bold=True, color=INK)
    note = cell.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(0)
    _add_run(note, hint, size=9.5, color=MUTED)

    anchor._p.addprevious(table._tbl)
    caption_paragraph = anchor.insert_paragraph_before()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(2)
    caption_paragraph.paragraph_format.space_after = Pt(5)
    _add_run(caption_paragraph, caption, size=10.5, color=MUTED)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _mark_row_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _restore_preserved_parts(template: Path, output: Path) -> None:
    editable_parts = {
        "[Content_Types].xml",
        "word/document.xml",
        "word/_rels/document.xml.rels",
    }
    with ZipFile(template) as source_archive:
        preserved = {
            info.filename: source_archive.read(info.filename)
            for info in source_archive.infolist()
            if info.filename not in editable_parts
        }

    temporary = output.with_name(f"{output.stem}.__tmp__.docx")
    try:
        with ZipFile(output) as current_archive, ZipFile(
            temporary,
            "w",
            compression=ZIP_DEFLATED,
        ) as rebuilt_archive:
            for info in current_archive.infolist():
                payload = preserved.get(info.filename)
                if payload is None:
                    payload = current_archive.read(info.filename)
                rebuilt_archive.writestr(info, payload)
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()


if __name__ == "__main__":
    main()

